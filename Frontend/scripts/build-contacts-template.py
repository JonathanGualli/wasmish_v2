"""
Genera Frontend/public/plantilla-contactos.xlsx, la plantilla que se descarga
en «Importar contactos». El archivo resultante va en el repo: esto solo hace
falta para cambiarla.

    pip install openpyxl
    python3 Frontend/scripts/build-contacts-template.py

La columna Teléfono va en formato texto: si no, Excel convierte los números
largos en 5,93991E+11 (y pierde dígitos) y les quita el 0 inicial.
Los nombres de las columnas tienen que coincidir con los que reconoce el
importador (guessColumnMapping).
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = [  # (cabecera, ancho)
    ("Teléfono", 20), ("Nombre", 18), ("Apellido", 18), ("Email", 28),
    ("Empresa", 24), ("Etiquetas", 24), ("Notas", 40),
]
ROWS_AS_TEXT = 10_001  # cabecera + el máximo de filas que admite el importador

INSTRUCTIONS = [
    ("Cómo llenar la plantilla", None),
    ("Una fila por contacto, en la hoja «Contactos». No cambies los nombres de las columnas.", None),
    (None, None),
    ("Columna", "Qué poner"),
    ("Teléfono", "Obligatorio. Con código de país (593991234567, +593 99 123 4567) o sin él (0991234567): "
                 "al importar eliges el país de los que no lo traen."),
    ("Nombre", "Opcional."),
    ("Apellido", "Opcional. Se junta con el nombre: «Ana» y «Ruiz» quedan «Ana Ruiz»."),
    ("Email", "Opcional."),
    ("Empresa", "Opcional."),
    ("Etiquetas", "Opcional. Separadas por comas: «VIP, Quito». Sirven para elegir a quién mandar una campaña."),
    ("Notas", "Opcional. Solo las ves tú."),
    (None, None),
    ("Si un teléfono ya está en Wasmish, solo se completan los datos que le falten: no se cambia nada de lo que ya tiene.", None),
    ("Puedes importar el mismo archivo otra vez: no se duplica ningún contacto.", None),
]

HEADER_FILL = PatternFill("solid", fgColor="EEF1F4")  # gris frío 100 del manual


def build(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Contactos"
    for index, (title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        sheet.column_dimensions[cell.column_letter].width = width
    for row in range(2, ROWS_AS_TEXT + 1):
        sheet.cell(row=row, column=1).number_format = "@"
    sheet.freeze_panes = "A2"

    help_sheet = book.create_sheet("Instrucciones")
    help_sheet.column_dimensions["A"].width = 16
    help_sheet.column_dimensions["B"].width = 90
    for row, (first, second) in enumerate(INSTRUCTIONS, start=1):
        if first:
            help_sheet.cell(row=row, column=1, value=first)
        if second:
            help_sheet.cell(row=row, column=2, value=second).alignment = Alignment(wrap_text=True, vertical="top")
    help_sheet["A1"].font = Font(bold=True, size=14)
    for cell in help_sheet[4]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    book.save(path)


if __name__ == "__main__":
    target = Path(__file__).resolve().parent.parent / "public" / "plantilla-contactos.xlsx"
    build(target)
    print(f"Escrita {target}")
